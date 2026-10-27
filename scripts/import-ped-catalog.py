import json,math,hashlib,argparse
from pathlib import Path
from openpyxl import load_workbook
parser=argparse.ArgumentParser(description='Import the supplied Appendix 19 workbook without inventing model powers')
parser.add_argument('source',type=Path)
args=parser.parse_args()
root=Path(__file__).resolve().parent.parent
src=args.source
old=json.loads((root/'assets/PedCoolingDatabase.json').read_text())
geometry=set((r['casing'],r['wall']) for r in old)
geometry.update((139.7,w) for c,w in list(geometry) if c==140)
models={}; omitted=[];generic=[]
s=load_workbook(src,data_only=True)['Таблица']
for row in s.iter_rows(min_row=10):
 m,n,d,p,v=[row[i].value for i in [3,4,5,6,7]]
 if not m:continue
 if n is None or not isinstance(d,(int,float)) or p is None:
  target=generic if d=='все' and p=='все' else omitted
  target.append(dict(manufacturer=m,name=n,diameter=d,power=p,speed=v,sourceRow=row[0].row));continue
 n=n.strip();p=p.strip() if isinstance(p,str) else p
 key=(m,n,d,p)
 model=dict(manufacturer=m,name=n,diameter=d,power=p,speed=v,sourceRow=row[0].row)
 if key in models:assert models[key]['speed']==v
 models[key]=model
records=[]
for model in models.values():
 for c,w in sorted(geometry):
  inner=c-2*w
  q=math.pi/4*((inner/1000)**2-(model['diameter']/1000)**2)*model['speed']*86400 if inner>model['diameter'] else None
  records.append(dict(model,casing=c,wall=w,flow=q))
meta=dict(source='Приложение 19, версия 7',sha256=hashlib.sha256(src.read_bytes()).hexdigest(),formula='PI()/4*((Dэк-2*s)^2-Dпэд^2)*1e-6*v*86400',models=list(models.values()),generic=generic,incomplete=omitted,geometries=[dict(casing=c,wall=w) for c,w in sorted(geometry)])
(root/'assets/PedCoolingCatalog-v7.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2)+'\n')
(root/'assets/PedCoolingDatabase.json').write_text(json.dumps(records,ensure_ascii=False,separators=(',',':'))+'\n')
print('models',len(models),'records',len(records),'valid',sum(r['flow'] is not None for r in records),'ranges',sorted({r['power'] for r in records if isinstance(r['power'],str)}),'generic',generic,'incomplete',omitted)
